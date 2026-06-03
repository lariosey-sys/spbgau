import os
import io
import csv
import json
import uuid
import time
import sqlite3
import hashlib
import threading
from datetime import datetime, timezone, timedelta
from functools import wraps
from flask import (
    Flask, render_template_string, jsonify, request,
    session, redirect, url_for, Response,
)
import paho.mqtt.client as mqtt

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "greenhouse-secret-key-2026")

MQTT_HOST = os.environ.get("MQTT_HOST", "mosquitto")
MQTT_PORT = int(os.environ.get("MQTT_PORT", "1883"))
RULES_FILE = os.environ.get("RULES_FILE", "/data/rules.json")
NAMES_FILE = os.environ.get("NAMES_FILE", "/data/names.json")
DB_FILE = os.environ.get("DB_FILE", "/data/history.db")
TZ_OFFSET = int(os.environ.get("TZ_OFFSET", "3"))  # MSK = UTC+3
PASSWORD = os.environ.get("DASHBOARD_PASSWORD", "7777")

DEFAULT_RELAY_NAMES = {
    1: "Бактерицидный фильтр", 2: "2-й УФ фильтр", 3: "1-й УФ фильтр",
    4: "1-й насос", 5: "4-й насос", 6: "2-й насос", 7: "3-й насос",
    8: "2-й вентилятор", 9: "4-я заслонка справа", 10: "1-й вентилятор",
    11: "3-я заслонка справа", 12: "2-я заслонка слева", 13: "Правый ТЭН",
    14: "1-я заслонка слева", 15: "Левый ТЭН",
}

DEFAULT_GROUP_NAMES = {
    "Заслонки": "Заслонки", "Вентиляторы": "Вентиляторы",
    "ТЭНы": "ТЭНы", "Насосы": "Насосы", "Фильтры": "Фильтры",
}

RELAY_GROUPS_MAP = {
    "Заслонки": [14, 12, 11, 9],
    "Вентиляторы": [10, 8],
    "ТЭНы": [15, 13],
    "Насосы": [4, 6, 7, 5],
    "Фильтры": [3, 2, 1],
}

# Custom names (loaded from file, override defaults)
custom_names = {"relays": {}, "groups": {}}

def load_names():
    global custom_names
    try:
        with open(NAMES_FILE, "r") as f:
            custom_names = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        custom_names = {"relays": {}, "groups": {}}

def save_names():
    os.makedirs(os.path.dirname(NAMES_FILE), exist_ok=True)
    with open(NAMES_FILE, "w") as f:
        json.dump(custom_names, f, ensure_ascii=False, indent=2)

def get_relay_names():
    names = dict(DEFAULT_RELAY_NAMES)
    for k, v in custom_names.get("relays", {}).items():
        names[int(k)] = v
    return names

def get_relay_groups():
    group_names = dict(DEFAULT_GROUP_NAMES)
    group_names.update(custom_names.get("groups", {}))
    return {group_names.get(k, k): v for k, v in RELAY_GROUPS_MAP.items()}

state = {
    "relays": {i: False for i in range(1, 16)},
    "arduino_online": False,
    "uptime": 0, "rssi": 0, "wifi_rc": 0, "mqtt_rc": 0,
    "sensors": {},
}
lock = threading.Lock()
mqtt_client = None

# --- Ventilation profiles ---
PROFILES_FILE = os.environ.get("PROFILES_FILE", "/data/profiles.json")
profiles = []
profiles_lock = threading.Lock()
active_profiles = {}  # profile_id -> True

def load_profiles():
    global profiles
    try:
        with open(PROFILES_FILE, "r") as f:
            profiles = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        profiles = []

def save_profiles():
    os.makedirs(os.path.dirname(PROFILES_FILE), exist_ok=True)
    with open(PROFILES_FILE, "w") as f:
        json.dump(profiles, f, ensure_ascii=False, indent=2)

def activate_profile(profile):
    """Activate: open dampers, wait, start fans/heaters."""
    pid = profile["id"]
    # Step 1: pre-actions (dampers)
    for a in profile.get("pre", []):
        r = a.get("relay")
        s = a.get("state", "ON")
        if r and mqtt_client:
            mqtt_client.publish(f"greenhouse/relay/{r}/set", s)

    delay = profile.get("delay", 10)
    if delay > 0:
        time.sleep(delay)

    # Check still active (may have been stopped during delay)
    if not active_profiles.get(pid):
        return

    # Step 2: main actions (fans, heaters)
    for a in profile.get("main", []):
        r = a.get("relay")
        s = a.get("state", "ON")
        if r and mqtt_client:
            mqtt_client.publish(f"greenhouse/relay/{r}/set", s)

def deactivate_profile(profile):
    """Deactivate: stop fans/heaters, wait, close dampers."""
    # Step 1: stop main (fans, heaters)
    for a in profile.get("main", []):
        r = a.get("relay")
        s = "OFF" if a.get("state", "ON") == "ON" else "ON"
        if r and mqtt_client:
            mqtt_client.publish(f"greenhouse/relay/{r}/set", s)

    delay = profile.get("delay_off", profile.get("delay", 10))
    if delay > 0:
        time.sleep(delay)

    # Step 2: reverse pre-actions (close dampers)
    for a in profile.get("pre", []):
        r = a.get("relay")
        s = "OFF" if a.get("state", "ON") == "ON" else "ON"
        if r and mqtt_client:
            mqtt_client.publish(f"greenhouse/relay/{r}/set", s)

def toggle_profile_async(profile_id):
    """Toggle profile in background thread (has delays)."""
    with profiles_lock:
        profile = None
        for p in profiles:
            if p["id"] == profile_id:
                profile = p
                break
    if not profile:
        return

    was_active = active_profiles.get(profile_id, False)
    if was_active:
        active_profiles[profile_id] = False
        threading.Thread(target=deactivate_profile, args=(profile,), daemon=True).start()
    else:
        active_profiles[profile_id] = True
        threading.Thread(target=activate_profile, args=(profile,), daemon=True).start()

# --- Auth ---
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated

# --- History DB ---
def init_db():
    os.makedirs(os.path.dirname(DB_FILE), exist_ok=True)
    conn = sqlite3.connect(DB_FILE)
    conn.execute("""CREATE TABLE IF NOT EXISTS sensor_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ts TEXT NOT NULL,
        device TEXT NOT NULL,
        temperature REAL,
        humidity REAL,
        co2 REAL
    )""")
    conn.execute("""CREATE TABLE IF NOT EXISTS relay_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ts TEXT NOT NULL,
        relay_id INTEGER NOT NULL,
        state TEXT NOT NULL
    )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_sensor_ts ON sensor_log(ts)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_relay_ts ON relay_log(ts)")
    conn.commit()
    conn.close()

def log_sensors():
    """Log current sensor readings to DB."""
    with lock:
        sensors = {k: dict(v) for k, v in state["sensors"].items()}
    now = now_local().strftime("%Y-%m-%d %H:%M:%S")
    conn = sqlite3.connect(DB_FILE)
    for dev, s in sensors.items():
        if not s.get("online", False):
            continue
        t = s.get("t")
        h = s.get("h")
        co2 = s.get("co2")
        if t is not None or h is not None or co2 is not None:
            conn.execute(
                "INSERT INTO sensor_log (ts, device, temperature, humidity, co2) VALUES (?,?,?,?,?)",
                (now, dev, t, h, co2),
            )
    conn.commit()
    conn.close()

def log_relay_change(relay_id, new_state):
    """Log relay state change."""
    now = now_local().strftime("%Y-%m-%d %H:%M:%S")
    conn = sqlite3.connect(DB_FILE)
    conn.execute(
        "INSERT INTO relay_log (ts, relay_id, state) VALUES (?,?,?)",
        (now, relay_id, new_state),
    )
    conn.commit()
    conn.close()

def history_logging_loop():
    """Background thread: log sensor data every 60 seconds."""
    while True:
        time.sleep(60)
        try:
            log_sensors()
        except Exception as e:
            print(f"[DB] Log error: {e}")

# --- Rules engine ---
rules = []
rules_lock = threading.Lock()
rule_active = {}

def now_local():
    return datetime.now(timezone(timedelta(hours=TZ_OFFSET)))

def load_rules():
    global rules
    try:
        with open(RULES_FILE, "r") as f:
            rules = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        rules = []

def save_rules():
    os.makedirs(os.path.dirname(RULES_FILE), exist_ok=True)
    with open(RULES_FILE, "w") as f:
        json.dump(rules, f, ensure_ascii=False, indent=2)

def check_condition(cond, sensors):
    device = cond.get("device", "")
    field = cond.get("field", "t")
    op = cond.get("op", ">")
    value = cond.get("value", 0)
    sensor = sensors.get(device)
    if not sensor or not sensor.get("online", False):
        return None
    reading = sensor.get(field)
    if reading is None:
        return None
    if op == ">": return reading > value
    elif op == "<": return reading < value
    elif op == ">=": return reading >= value
    elif op == "<=": return reading <= value
    elif op == "==": return abs(reading - value) < 0.01
    return None

def check_schedule(schedule):
    if not schedule or not schedule.get("enabled", False):
        return True
    now = now_local()
    current_day = now.weekday()
    days = schedule.get("days", [0,1,2,3,4,5,6])
    if current_day not in days:
        return False
    current_minutes = now.hour * 60 + now.minute
    start = schedule.get("start_minutes", 0)
    end = schedule.get("end_minutes", 1440)
    if start <= end:
        return start <= current_minutes < end
    else:
        return current_minutes >= start or current_minutes < end

def evaluate_rule_with_hysteresis(rule, sensors, was_active):
    if not rule.get("enabled", True):
        return False
    if not check_schedule(rule.get("schedule")):
        return False
    conditions = rule.get("conditions", [])
    if not conditions:
        return True
    for cond in conditions:
        hysteresis = cond.get("hysteresis", 0)
        device = cond.get("device", "")
        field = cond.get("field", "t")
        op = cond.get("op", ">")
        value = cond.get("value", 0)
        sensor = sensors.get(device)
        if not sensor or not sensor.get("online", False):
            return False
        reading = sensor.get(field)
        if reading is None:
            return False
        if was_active and hysteresis > 0:
            if op in (">", ">=") and reading < value - hysteresis:
                return False
            elif op in ("<", "<=") and reading > value + hysteresis:
                return False
        else:
            result = check_condition(cond, sensors)
            if not result:
                return False
    return True

def execute_actions(actions, reverse=False):
    if not mqtt_client:
        return
    for action in actions:
        relay = action.get("relay")
        command = action.get("command", "ON")
        if reverse:
            command = "OFF" if command == "ON" else "ON"
        if relay and 1 <= relay <= 15:
            mqtt_client.publish(f"greenhouse/relay/{relay}/set", command)

def find_profile(pid):
    with profiles_lock:
        for p in profiles:
            if p["id"] == pid:
                return dict(p)
    return None

def execute_rule_action(rule, activate):
    """Execute rule action: either direct relays or a profile."""
    profile_id = rule.get("profile_id")
    if profile_id:
        profile = find_profile(profile_id)
        if not profile:
            return
        if activate:
            active_profiles[profile_id] = True
            threading.Thread(target=activate_profile, args=(profile,), daemon=True).start()
        else:
            active_profiles[profile_id] = False
            threading.Thread(target=deactivate_profile, args=(profile,), daemon=True).start()
    else:
        if activate:
            execute_actions(rule.get("actions", []))
        else:
            execute_actions(rule.get("actions", []), reverse=True)

def rules_evaluation_loop():
    while True:
        time.sleep(5)
        with lock:
            sensors = {k: dict(v) for k, v in state["sensors"].items()}
        with rules_lock:
            for rule in rules:
                if not rule.get("enabled", True):
                    continue
                rid = rule.get("id", "")
                was_active = rule_active.get(rid, False)
                now_active = evaluate_rule_with_hysteresis(rule, sensors, was_active)
                if now_active and not was_active:
                    execute_rule_action(rule, True)
                    rule_active[rid] = True
                elif not now_active and was_active and rule.get("reverse", True):
                    execute_rule_action(rule, False)
                    rule_active[rid] = False
                elif not now_active:
                    rule_active[rid] = False

# --- MQTT ---
def on_connect(client, userdata, flags, rc, properties=None):
    client.subscribe("greenhouse/relay/+/state")
    client.subscribe("greenhouse/mega-1/status")
    client.subscribe("greenhouse/relay/summary")
    client.subscribe("greenhouse/env/+/state")
    client.subscribe("greenhouse/+/status")

def on_message(client, userdata, msg):
    topic = msg.topic
    payload = msg.payload.decode()
    with lock:
        if topic == "greenhouse/mega-1/status":
            state["arduino_online"] = payload == "online"
        elif topic == "greenhouse/relay/summary":
            try:
                data = json.loads(payload)
                state["uptime"] = data.get("uptime", 0)
                state["rssi"] = data.get("rssi", 0)
                state["wifi_rc"] = data.get("wifi_rc", 0)
                state["mqtt_rc"] = data.get("mqtt_rc", 0)
                for k, v in data.get("relays", {}).items():
                    state["relays"][int(k)] = bool(v)
            except (json.JSONDecodeError, ValueError):
                pass
        elif topic.startswith("greenhouse/env/"):
            try:
                data = json.loads(payload)
                device = data.get("device", "")
                if device:
                    s = state["sensors"].setdefault(device, {})
                    s["t"] = data.get("t")
                    s["h"] = data.get("h")
                    s["co2"] = data.get("co2")
                    s["ts"] = data.get("ts", 0)
            except (json.JSONDecodeError, ValueError):
                pass
        elif topic.endswith("/status"):
            device = topic.split("/")[1]
            if device.startswith(("th-", "co2-")):
                state["sensors"].setdefault(device, {})["online"] = payload == "online"
        elif "/state" in topic:
            try:
                num = int(topic.split("/")[-2])
                old = state["relays"].get(num, False)
                new = payload == "ON"
                state["relays"][num] = new
                if old != new:
                    try:
                        log_relay_change(num, payload)
                    except Exception:
                        pass
            except (ValueError, IndexError):
                pass

def start_mqtt():
    global mqtt_client
    mqtt_client = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2)
    mqtt_client.on_connect = on_connect
    mqtt_client.on_message = on_message
    mqtt_client.connect(MQTT_HOST, MQTT_PORT, 60)
    mqtt_client.loop_start()


# ===================== TEMPLATES =====================

STYLE = """
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    background: #0f1923; color: #e0e0e0; min-height: 100vh; padding: 16px;
  }
  .header {
    display: flex; align-items: center; justify-content: space-between;
    padding: 16px 20px; background: #1a2733; border-radius: 12px;
    margin-bottom: 20px; flex-wrap: wrap; gap: 10px;
  }
  .header h1 { font-size: 1.4em; color: #4fc3f7; }
  .status-bar { display: flex; gap: 16px; align-items: center; flex-wrap: wrap; }
  .status-item { font-size: 0.85em; color: #90a4ae; }
  .status-item span { color: #e0e0e0; font-weight: 600; }
  .dot {
    display: inline-block; width: 10px; height: 10px; border-radius: 50%;
    margin-right: 6px; vertical-align: middle;
  }
  .dot.online { background: #4caf50; box-shadow: 0 0 8px #4caf5088; }
  .dot.offline { background: #f44336; box-shadow: 0 0 8px #4caf5088; }
  .tabs { display: flex; gap: 8px; margin-bottom: 20px; }
  .tab {
    padding: 10px 24px; border-radius: 8px; border: 1px solid #37474f;
    background: transparent; color: #90a4ae; cursor: pointer;
    font-size: 0.95em; font-weight: 600; text-decoration: none;
  }
  .tab:hover { background: #263238; color: #e0e0e0; }
  .tab.active { background: #1a2733; color: #4fc3f7; border-color: #4fc3f7; }
  .group {
    background: #1a2733; border-radius: 12px; padding: 16px 20px; margin-bottom: 16px;
  }
  .group-header {
    display: flex; align-items: center; justify-content: space-between; margin-bottom: 12px;
  }
  .group-title { font-size: 1.1em; color: #4fc3f7; font-weight: 600; }
  .group-actions { display: flex; gap: 8px; }
  .group-btn {
    padding: 4px 12px; border-radius: 6px; border: 1px solid #37474f;
    background: transparent; color: #90a4ae; cursor: pointer; font-size: 0.8em;
  }
  .group-btn:hover { background: #263238; color: #e0e0e0; }
  .relays { display: grid; grid-template-columns: repeat(auto-fill, minmax(200px, 1fr)); gap: 10px; }
  .relay {
    display: flex; align-items: center; justify-content: space-between;
    padding: 12px 16px; background: #263238; border-radius: 8px; transition: background 0.2s;
  }
  .relay.on { background: #1b3a2a; }
  .relay-name { font-size: 0.95em; cursor: pointer; }
  .relay-name:hover { color: #4fc3f7; }
  .rename-input {
    background: #1a2733; border: 1px solid #4fc3f7; border-radius: 4px;
    color: #e0e0e0; padding: 4px 8px; font-size: 0.9em; width: 140px;
  }
  .group-title-edit:hover { cursor: pointer; color: #81d4fa; }
  .toggle { position: relative; width: 48px; height: 26px; cursor: pointer; }
  .toggle input { display: none; }
  .toggle .slider {
    position: absolute; inset: 0; background: #37474f;
    border-radius: 13px; transition: background 0.3s;
  }
  .toggle .slider::before {
    content: ''; position: absolute; width: 20px; height: 20px;
    left: 3px; top: 3px; background: #ccc; border-radius: 50%;
    transition: transform 0.3s, background 0.3s;
  }
  .toggle input:checked + .slider { background: #4caf50; }
  .toggle input:checked + .slider::before { transform: translateX(22px); background: #fff; }
  .sensors {
    display: grid; grid-template-columns: repeat(auto-fill, minmax(300px, 1fr));
    gap: 12px; margin-bottom: 20px;
  }
  .sensor-card { background: #1a2733; border-radius: 12px; padding: 16px 20px; }
  .sensor-header {
    display: flex; align-items: center; justify-content: space-between; margin-bottom: 10px;
  }
  .sensor-name { font-size: 0.9em; color: #90a4ae; }
  .sensor-values { display: flex; gap: 16px; flex-wrap: wrap; }
  .sensor-val { font-size: 1.6em; font-weight: 700; line-height: 1.2; }
  .sensor-val small { font-size: 0.45em; color: #90a4ae; display: block; font-weight: 400; }
  .temp-color { color: #ff9800; }
  .hum-color { color: #29b6f6; }
  .co2-color { color: #66bb6a; }
  .sensor-offline { opacity: 0.4; }
  .all-controls { display: flex; gap: 10px; margin-bottom: 20px; }
  .all-btn {
    padding: 10px 24px; border-radius: 8px; border: none;
    font-size: 0.95em; cursor: pointer; font-weight: 600;
  }
  .all-btn.on-btn { background: #4caf50; color: #fff; }
  .all-btn.on-btn:hover { background: #43a047; }
  .all-btn.off-btn { background: #f44336; color: #fff; }
  .all-btn.off-btn:hover { background: #e53935; }
  .rule-card {
    background: #1a2733; border-radius: 12px; padding: 16px 20px;
    margin-bottom: 12px; border-left: 4px solid #37474f;
  }
  .rule-card.active { border-left-color: #4caf50; }
  .rule-card.disabled { opacity: 0.5; }
  .rule-header {
    display: flex; align-items: center; justify-content: space-between; margin-bottom: 8px;
  }
  .rule-name { font-size: 1.05em; font-weight: 600; }
  .rule-actions { display: flex; gap: 8px; align-items: center; }
  .rule-desc { font-size: 0.85em; color: #90a4ae; line-height: 1.5; }
  .rule-status {
    display: inline-block; padding: 2px 8px; border-radius: 4px;
    font-size: 0.75em; font-weight: 600;
  }
  .rule-status.active { background: #1b3a2a; color: #4caf50; }
  .rule-status.inactive { background: #263238; color: #90a4ae; }
  .btn {
    padding: 6px 14px; border-radius: 6px; border: 1px solid #37474f;
    background: transparent; color: #90a4ae; cursor: pointer; font-size: 0.85em;
  }
  .btn:hover { background: #263238; color: #e0e0e0; }
  .btn.danger { border-color: #f44336; color: #f44336; }
  .btn.danger:hover { background: #3a1a1a; }
  .btn.primary { border-color: #4fc3f7; color: #4fc3f7; }
  .btn.primary:hover { background: #1a2f3a; }
  .form-card {
    background: #1a2733; border-radius: 12px; padding: 20px; margin-bottom: 20px;
  }
  .form-card h2 { font-size: 1.1em; color: #4fc3f7; margin-bottom: 16px; }
  .form-row {
    display: flex; gap: 12px; margin-bottom: 12px; align-items: end; flex-wrap: wrap;
  }
  .form-group { display: flex; flex-direction: column; gap: 4px; }
  .form-group label { font-size: 0.8em; color: #90a4ae; }
  .form-group input, .form-group select {
    padding: 8px 12px; border-radius: 6px; border: 1px solid #37474f;
    background: #263238; color: #e0e0e0; font-size: 0.9em; min-width: 0;
  }
  .form-group input:focus, .form-group select:focus { outline: none; border-color: #4fc3f7; }
  .checkbox-row { display: flex; gap: 12px; flex-wrap: wrap; align-items: center; }
  .checkbox-row label {
    font-size: 0.85em; display: flex; align-items: center; gap: 4px; cursor: pointer;
  }
  .checkbox-row input[type="checkbox"] { accent-color: #4fc3f7; }
  .section-label { font-size: 0.9em; color: #4fc3f7; margin: 16px 0 8px 0; font-weight: 600; }
  .relay-checkboxes { display: flex; flex-wrap: wrap; gap: 8px; }
  .relay-chip {
    display: flex; align-items: center; gap: 4px; padding: 4px 10px;
    border-radius: 6px; background: #263238; font-size: 0.8em; cursor: pointer;
  }
  .relay-chip input { accent-color: #4caf50; }
  .profiles { display: grid; grid-template-columns: repeat(auto-fill, minmax(260px, 1fr)); gap: 12px; margin-bottom: 20px; }
  .profile-card {
    background: #1a2733; border-radius: 12px; padding: 16px; border: 2px solid #263238;
    cursor: pointer; transition: all 0.2s;
  }
  .profile-card:hover { border-color: #37474f; }
  .profile-card.active { border-color: #4caf50; background: #1b3a2a; }
  .profile-name { font-weight: 600; margin-bottom: 6px; }
  .profile-desc { font-size: 0.8em; color: #90a4ae; line-height: 1.4; }
  .profile-badge {
    display: inline-block; padding: 2px 8px; border-radius: 4px;
    font-size: 0.7em; font-weight: 600; margin-top: 8px;
  }
  .profile-badge.on { background: #1b3a2a; color: #4caf50; }
  .profile-badge.off { background: #263238; color: #607d8b; }
  .login-box {
    max-width: 320px; margin: 80px auto; background: #1a2733;
    border-radius: 12px; padding: 32px;
  }
  .login-box h1 { color: #4fc3f7; margin-bottom: 20px; font-size: 1.3em; text-align: center; }
  .login-box input[type=password] {
    width: 100%; padding: 12px; border-radius: 8px; border: 1px solid #37474f;
    background: #263238; color: #e0e0e0; font-size: 1em; margin-bottom: 16px;
  }
  .login-box button {
    width: 100%; padding: 12px; border-radius: 8px; border: none;
    background: #4fc3f7; color: #0f1923; font-size: 1em; font-weight: 600; cursor: pointer;
  }
  .login-box .error { color: #f44336; font-size: 0.85em; margin-bottom: 12px; }
  .stats-table {
    width: 100%; border-collapse: collapse; margin-top: 12px;
  }
  .stats-table th, .stats-table td {
    padding: 8px 12px; text-align: left; border-bottom: 1px solid #263238;
    font-size: 0.85em;
  }
  .stats-table th { color: #4fc3f7; font-weight: 600; }
  .stats-table tr:hover { background: #1a2733; }
  .export-buttons { display: flex; gap: 8px; margin-top: 16px; }
  .export-btn {
    padding: 8px 20px; border-radius: 6px; border: none;
    font-size: 0.9em; cursor: pointer; font-weight: 600; text-decoration: none;
  }
  .export-btn.csv { background: #ff9800; color: #fff; }
  .export-btn.excel { background: #4caf50; color: #fff; }
"""

TEMPLATE = """<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Теплица</title>
<script src="https://unpkg.com/htmx.org@2.0.4"></script>
<style>""" + STYLE + """</style>
</head>
<body>
{{ content | safe }}
<script>
function startRename(el, type, id) {
  if (el.querySelector('input')) return;
  var old = el.textContent.trim();
  el.innerHTML = '<input class="rename-input" value="' + old + '" autofocus>';
  var inp = el.querySelector('input');
  inp.focus(); inp.select();
  function save() {
    var val = inp.value.trim();
    if (val && val !== old) {
      fetch('/api/rename', {method:'POST', headers:{'Content-Type':'application/json'},
        body: JSON.stringify({type:type, id:id, name:val})
      }).then(function() { el.textContent = val; });
    } else { el.textContent = old; }
  }
  inp.addEventListener('blur', save);
  inp.addEventListener('keydown', function(e) {
    if (e.key === 'Enter') { e.preventDefault(); inp.blur(); }
    if (e.key === 'Escape') { el.textContent = old; }
  });
}
</script>
</body>
</html>"""

LOGIN_PAGE = """
<div class="login-box">
  <h1>Теплица</h1>
  {% if error %}<div class="error">{{ error }}</div>{% endif %}
  <form method="POST">
    <input type="password" name="password" placeholder="Пароль" autofocus>
    <button type="submit">Войти</button>
  </form>
</div>
"""

SENSORS_BLOCK = """
{% if sensors %}
<div class="sensors">
  {% for dev, s in sensors.items() %}
  {% if s.get('online', False) or s.get('t') is not none or s.get('h') is not none %}
  <div class="sensor-card {{ '' if s.get('online', False) else 'sensor-offline' }}">
    <div class="sensor-header">
      <span class="sensor-name">{{ dev }}</span>
      <span class="dot {{ 'online' if s.get('online', False) else 'offline' }}"></span>
    </div>
    <div class="sensor-values">
      {% if s.get('t') is not none %}
      <div class="sensor-val temp-color">{{ "%.1f"|format(s.t) }}°C<small>Температура</small></div>
      {% endif %}
      {% if s.get('h') is not none %}
      <div class="sensor-val hum-color">{{ "%.1f"|format(s.h) }}%<small>Влажность</small></div>
      {% endif %}
      {% if s.get('co2') is not none %}
      <div class="sensor-val co2-color">{{ "%.0f"|format(s.co2) }}<small>CO2 ppm</small></div>
      {% endif %}
    </div>
  </div>
  {% endif %}
  {% endfor %}
</div>
{% endif %}
"""

STATE_FRAGMENT = """
<div class="header">
  <h1>Теплица</h1>
  <div class="status-bar">
    <div class="status-item">
      <span class="dot {{ 'online' if arduino_online else 'offline' }}"></span>
      Arduino {{ 'Online' if arduino_online else 'Offline' }}
    </div>
    <div class="status-item">Uptime: <span>{{ uptime_str }}</span></div>
    <div class="status-item">RSSI: <span>{{ rssi }}</span></div>
  </div>
</div>

<div class="tabs">
  <a class="tab active" href="/">Управление</a>
  <a class="tab" href="/rules">Правила</a>
  <a class="tab" href="/profiles">Вентиляция</a>
  <a class="tab" href="/stats">Статистика</a>
</div>

""" + SENSORS_BLOCK + """

{% if profiles %}
<div class="group">
  <div class="group-header">
    <div class="group-title">Режимы вентиляции</div>
    <div class="group-actions">
      <a class="group-btn" href="/profiles">Настроить</a>
    </div>
  </div>
  <div class="profiles">
    {% for p in profiles %}
    <div class="profile-card {{ 'active' if active_profiles.get(p.id) else '' }}"
         hx-post="/api/profile/{{ p.id }}/toggle" hx-swap="none"
         onclick="this.classList.toggle('active')">
      <div class="profile-name">{{ p.name }}</div>
      <div class="profile-desc">{{ p.get('description', '') }}</div>
      <span class="profile-badge {{ 'on' if active_profiles.get(p.id) else 'off' }}">
        {{ 'Работает' if active_profiles.get(p.id) else 'Выкл' }}
      </span>
    </div>
    {% endfor %}
  </div>
</div>
{% else %}
<div class="group">
  <div class="group-header">
    <div class="group-title">Режимы вентиляции</div>
    <div class="group-actions"><a class="group-btn" href="/profiles">Настроить</a></div>
  </div>
  <div style="color:#607d8b; padding:8px 0; font-size:0.9em;">Нет настроенных режимов. <a href="/profiles" style="color:#4fc3f7;">Создать</a></div>
</div>
{% endif %}

<div class="all-controls">
  <button class="all-btn on-btn" hx-post="/relay/all/on" hx-swap="none">Включить всё</button>
  <button class="all-btn off-btn" hx-post="/relay/all/off" hx-swap="none">Выключить всё</button>
</div>

{% for group_name, relay_ids in groups.items() %}
<div class="group">
  <div class="group-header">
    <div class="group-title group-title-edit" onclick="startRename(this,'group','{{ group_name }}')">{{ group_name }}</div>
    <div class="group-actions">
      <button class="group-btn" hx-post="/group/{{ group_name }}/on" hx-swap="none">Вкл</button>
      <button class="group-btn" hx-post="/group/{{ group_name }}/off" hx-swap="none">Выкл</button>
    </div>
  </div>
  <div class="relays">
    {% for rid in relay_ids %}
    <div class="relay {{ 'on' if relays[rid] else '' }}">
      <span class="relay-name" onclick="startRename(this,'relay','{{ rid }}')">{{ names[rid] }}</span>
      <label class="toggle">
        <input type="checkbox" {{ 'checked' if relays[rid] else '' }}
               hx-post="/relay/{{ rid }}/toggle" hx-swap="none">
        <span class="slider"></span>
      </label>
    </div>
    {% endfor %}
  </div>
</div>
{% endfor %}
"""

RULES_PAGE = """
<div class="header">
  <h1>Теплица</h1>
  <div class="status-bar">
    <div class="status-item">
      <span class="dot {{ 'online' if arduino_online else 'offline' }}"></span>
      Arduino {{ 'Online' if arduino_online else 'Offline' }}
    </div>
  </div>
</div>

<div class="tabs">
  <a class="tab" href="/">Управление</a>
  <a class="tab active" href="/rules">Правила</a>
  <a class="tab" href="/profiles">Вентиляция</a>
  <a class="tab" href="/stats">Статистика</a>
</div>

""" + SENSORS_BLOCK + """

<div class="form-card" id="rule-form">
  <h2>Новое правило</h2>
  <form hx-post="/api/rules" hx-target="#rules-list" hx-swap="innerHTML" hx-on::after-request="if(event.detail.successful) this.reset()">
    <input type="hidden" name="edit_id" value="">
    <div class="form-row">
      <div class="form-group" style="flex:1; min-width:200px;">
        <label>Название</label>
        <input type="text" name="name" required placeholder="Вентиляция при жаре">
      </div>
    </div>
    <div class="section-label">Условие по датчику (необязательно)</div>
    <div class="form-row">
      <div class="form-group"><label>Датчик</label>
        <select name="cond_device"><option value="">-- нет --</option>
        {% for dev in sensor_list %}<option value="{{ dev }}">{{ dev }}</option>{% endfor %}
        </select>
      </div>
      <div class="form-group"><label>Параметр</label>
        <select name="cond_field">
          <option value="t">Температура</option><option value="h">Влажность</option>
          <option value="co2">CO2 (ppm)</option>
        </select>
      </div>
      <div class="form-group"><label>Условие</label>
        <select name="cond_op">
          <option value=">">></option><option value="<"><</option>
          <option value=">=">>=</option><option value="<="><=</option>
        </select>
      </div>
      <div class="form-group"><label>Значение</label>
        <input type="number" name="cond_value" step="0.1" placeholder="30" style="width:80px;">
      </div>
      <div class="form-group"><label>Гистерезис</label>
        <input type="number" name="cond_hysteresis" step="0.1" value="2" style="width:80px;">
      </div>
    </div>
    <div class="section-label">Расписание (необязательно)</div>
    <div class="form-row">
      <div class="form-group"><label>С</label><input type="time" name="sched_start"></div>
      <div class="form-group"><label>По</label><input type="time" name="sched_end"></div>
    </div>
    <div class="checkbox-row" style="margin-bottom:12px;">
      <label><input type="checkbox" name="day_0" checked> Пн</label>
      <label><input type="checkbox" name="day_1" checked> Вт</label>
      <label><input type="checkbox" name="day_2" checked> Ср</label>
      <label><input type="checkbox" name="day_3" checked> Чт</label>
      <label><input type="checkbox" name="day_4" checked> Пт</label>
      <label><input type="checkbox" name="day_5" checked> Сб</label>
      <label><input type="checkbox" name="day_6" checked> Вс</label>
    </div>
    <div class="section-label">Действие</div>
    <div class="form-row">
      <div class="form-group"><label>Режим вентиляции (приоритет)</label>
        <select name="profile_id">
          <option value="">-- не использовать --</option>
          {% for p in profiles_list %}
          <option value="{{ p.id }}">{{ p.name }}</option>
          {% endfor %}
        </select>
      </div>
    </div>
    <div class="section-label" style="color:#607d8b;">Или отдельные реле (если режим не выбран)</div>
    <div class="relay-checkboxes" style="margin-bottom:12px;">
      {% for rid, rname in relay_names.items()|sort %}
      <label class="relay-chip"><input type="checkbox" name="relay_{{ rid }}"> {{ rname }}</label>
      {% endfor %}
    </div>
    <div class="checkbox-row" style="margin-bottom:16px;">
      <label><input type="checkbox" name="reverse" checked> Откат при невыполнении условия</label>
    </div>
    <button type="submit" class="btn primary">Сохранить правило</button>
  </form>
</div>
<div id="rules-list">RULES_LIST_PLACEHOLDER</div>
"""

RULES_LIST_FRAGMENT = """
{% if rules %}
{% for rule in rules %}
<div class="rule-card {{ 'active' if active_map.get(rule.id) else '' }} {{ 'disabled' if not rule.enabled else '' }}">
  <div class="rule-header">
    <div>
      <span class="rule-name">{{ rule.name }}</span>
      {% if rule.enabled %}
      <span class="rule-status {{ 'active' if active_map.get(rule.id) else 'inactive' }}">
        {{ 'Работает' if active_map.get(rule.id) else 'Ожидает' }}
      </span>
      {% else %}
      <span class="rule-status inactive">Выключено</span>
      {% endif %}
    </div>
    <div class="rule-actions">
      <button class="btn" hx-post="/api/rules/{{ rule.id }}/toggle" hx-target="#rules-list" hx-swap="innerHTML">
        {{ 'Выкл' if rule.enabled else 'Вкл' }}
      </button>
      <button class="btn danger" hx-delete="/api/rules/{{ rule.id }}" hx-target="#rules-list" hx-swap="innerHTML"
              hx-confirm="Удалить правило '{{ rule.name }}'?">Удалить</button>
    </div>
  </div>
  <div class="rule-desc">
    {% for c in rule.get('conditions', []) %}
      {{ c.device }}: {{ {'t':'температура','h':'влажность','co2':'CO2'}.get(c.field, c.field) }} {{ c.op }} {{ c.value }}{{ {'t':'°C','h':'%','co2':' ppm'}.get(c.field, '') }}
      {% if c.hysteresis %} (гист. {{ c.hysteresis }}){% endif %}<br>
    {% endfor %}
    {% if rule.get('schedule', {}).get('enabled') %}
      {% set sch = rule.schedule %}
      Время: {{ "%02d:%02d"|format(sch.start_minutes // 60, sch.start_minutes % 60) }} - {{ "%02d:%02d"|format(sch.end_minutes // 60, sch.end_minutes % 60) }}
      {% set days = sch.get('days', []) %}
      {% if days|length == 7 %}(ежедневно){% endif %}
    {% endif %}
    {% if rule.get('profile_id') %}
      Режим: {{ rule.get('profile_name', rule.profile_id) }}
    {% else %}
      Реле: {% for a in rule.get('actions', []) %}{{ relay_names.get(a.relay, a.relay) }}{% if not loop.last %}, {% endif %}{% endfor %}
    {% endif %}
    {% if rule.get('reverse') %} (с откатом){% endif %}
  </div>
</div>
{% endfor %}
{% else %}
<div style="text-align:center; color:#90a4ae; padding:40px;">Правил пока нет.</div>
{% endif %}
"""

PROFILES_PAGE = """
<div class="header">
  <h1>Теплица</h1>
  <div class="status-bar">
    <div class="status-item">
      <span class="dot {{ 'online' if arduino_online else 'offline' }}"></span>
      Arduino {{ 'Online' if arduino_online else 'Offline' }}
    </div>
  </div>
</div>

<div class="tabs">
  <a class="tab" href="/">Управление</a>
  <a class="tab" href="/rules">Правила</a>
  <a class="tab active" href="/profiles">Вентиляция</a>
  <a class="tab" href="/stats">Статистика</a>
</div>

<div class="form-card">
  <h2>Новый режим вентиляции</h2>
  <form hx-post="/api/profiles" hx-target="#profiles-list" hx-swap="innerHTML" hx-on::after-request="if(event.detail.successful) this.reset()">
    <div class="form-row">
      <div class="form-group" style="flex:1; min-width:200px;">
        <label>Название</label>
        <input type="text" name="name" required placeholder="Приток с улицы (Бокс 1)">
      </div>
    </div>
    <div class="form-row">
      <div class="form-group" style="flex:1; min-width:200px;">
        <label>Описание</label>
        <input type="text" name="description" placeholder="Обе заслонки открыты, вентилятор включён">
      </div>
    </div>

    <div class="section-label">Шаг 1: Заслонки (выполняются первыми)</div>
    <div class="relay-checkboxes" style="margin-bottom:8px;">
      {% for rid in [14,12,11,9] %}
      <label class="relay-chip">
        <input type="checkbox" name="pre_{{ rid }}"> {{ relay_names.get(rid) }}
      </label>
      {% endfor %}
    </div>

    <div class="form-row">
      <div class="form-group">
        <label>Задержка перед шагом 2 (сек)</label>
        <input type="number" name="delay" value="10" style="width:80px;">
      </div>
      <div class="form-group">
        <label>Задержка при выключении (сек)</label>
        <input type="number" name="delay_off" value="10" style="width:80px;">
      </div>
    </div>

    <div class="section-label">Шаг 2: Вентиляторы и ТЭНы (после задержки)</div>
    <div class="relay-checkboxes" style="margin-bottom:8px;">
      {% for rid in [10,8] %}
      <label class="relay-chip">
        <input type="checkbox" name="main_{{ rid }}"> {{ relay_names.get(rid) }}
      </label>
      {% endfor %}
      {% for rid in [15,13] %}
      <label class="relay-chip">
        <input type="checkbox" name="main_{{ rid }}"> {{ relay_names.get(rid) }}
      </label>
      {% endfor %}
    </div>

    <button type="submit" class="btn primary" style="margin-top:12px;">Создать режим</button>
  </form>
</div>

<div id="profiles-list">
  PROFILES_LIST_PLACEHOLDER
</div>
"""

PROFILES_LIST_FRAGMENT = """
{% if profiles %}
{% for p in profiles %}
<div class="rule-card {{ 'active' if active_profiles.get(p.id) else '' }}">
  <div class="rule-header">
    <div>
      <span class="rule-name">{{ p.name }}</span>
      <span class="rule-status {{ 'active' if active_profiles.get(p.id) else 'inactive' }}">
        {{ 'Работает' if active_profiles.get(p.id) else 'Выкл' }}
      </span>
    </div>
    <div class="rule-actions">
      <button class="btn {{ 'danger' if active_profiles.get(p.id) else 'primary' }}"
              hx-post="/api/profile/{{ p.id }}/toggle" hx-target="#profiles-list" hx-swap="innerHTML">
        {{ 'Остановить' if active_profiles.get(p.id) else 'Запустить' }}
      </button>
      <button class="btn danger" hx-delete="/api/profile/{{ p.id }}" hx-target="#profiles-list" hx-swap="innerHTML"
              hx-confirm="Удалить режим '{{ p.name }}'?">Удалить</button>
    </div>
  </div>
  <div class="rule-desc">
    Заслонки: {% for a in p.get('pre', []) %}{{ relay_names.get(a.relay, a.relay) }}{% if not loop.last %}, {% endif %}{% endfor %}
    → задержка {{ p.get('delay', 10) }} сек →
    {% for a in p.get('main', []) %}{{ relay_names.get(a.relay, a.relay) }}{% if not loop.last %}, {% endif %}{% endfor %}
    <br>При выключении: обратный порядок (задержка {{ p.get('delay_off', p.get('delay', 10)) }} сек)
    {% if p.get('description') %}<br>{{ p.description }}{% endif %}
  </div>
</div>
{% endfor %}
{% else %}
<div style="text-align:center; color:#90a4ae; padding:40px;">
  Режимов пока нет. Создайте первый выше.
</div>
{% endif %}
"""

STATS_PAGE = """
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.7/dist/chart.umd.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/chartjs-adapter-date-fns@3.0.0/dist/chartjs-adapter-date-fns.bundle.min.js"></script>

<div class="header">
  <h1>Теплица</h1>
  <div class="status-bar">
    <div class="status-item">
      <span class="dot {{ 'online' if arduino_online else 'offline' }}"></span>
      Arduino {{ 'Online' if arduino_online else 'Offline' }}
    </div>
  </div>
</div>

<div class="tabs">
  <a class="tab" href="/">Управление</a>
  <a class="tab" href="/rules">Правила</a>
  <a class="tab" href="/profiles">Вентиляция</a>
  <a class="tab active" href="/stats">Статистика</a>
</div>

<div class="form-card">
  <h2>Графики</h2>
  <div class="form-row">
    <div class="form-group"><label>Период</label>
      <select id="chart_range">
        <option value="30">30 минут</option>
        <option value="60">1 час</option>
        <option value="180">3 часа</option>
        <option value="360" selected>6 часов</option>
        <option value="720">12 часов</option>
        <option value="1440">24 часа</option>
        <option value="4320">3 дня</option>
        <option value="10080">7 дней</option>
      </select>
    </div>
    <div class="form-group"><label>Обновление</label>
      <select id="chart_refresh">
        <option value="0">Выкл</option>
        <option value="10" selected>10 сек</option>
        <option value="30">30 сек</option>
        <option value="60">1 мин</option>
      </select>
    </div>
    <div class="form-group"><label>Датчики</label>
      <div class="checkbox-row">
        {% for dev in sensor_list %}
        <label><input type="checkbox" class="chart-device" value="{{ dev }}" checked> {{ dev }}</label>
        {% endfor %}
      </div>
    </div>
  </div>
  <div class="checkbox-row" style="margin-bottom:8px;">
    <label><input type="checkbox" id="show_temp" checked> Температура</label>
    <label><input type="checkbox" id="show_hum" checked> Влажность</label>
    <label><input type="checkbox" id="show_co2" checked> CO2</label>
  </div>
</div>

<div class="form-card" style="padding:12px 16px;">
  <h2 style="margin-bottom:8px;">Температура, °C</h2>
  <div style="position:relative; height:220px;"><canvas id="chartTemp"></canvas></div>
</div>
<div class="form-card" style="padding:12px 16px;">
  <h2 style="margin-bottom:8px;">Влажность, %</h2>
  <div style="position:relative; height:220px;"><canvas id="chartHum"></canvas></div>
</div>
<div class="form-card" style="padding:12px 16px;">
  <h2 style="margin-bottom:8px;">CO2, ppm</h2>
  <div style="position:relative; height:220px;"><canvas id="chartCO2"></canvas></div>
</div>

<div class="form-card">
  <h2>Выгрузка данных</h2>
  <form id="export-form">
    <div class="form-row">
      <div class="form-group"><label>Тип данных</label>
        <select id="data_type">
          <option value="sensors">Датчики (температура, влажность, CO2)</option>
          <option value="relays">Реле (вкл/выкл)</option>
        </select>
      </div>
      <div class="form-group"><label>Датчик / Реле</label>
        <select id="device_filter">
          <option value="">Все</option>
          {% for dev in sensor_list %}<option value="{{ dev }}">{{ dev }}</option>{% endfor %}
          {% for rid, rname in relay_names.items()|sort %}<option value="relay_{{ rid }}">{{ rname }}</option>{% endfor %}
        </select>
      </div>
    </div>
    <div class="form-row">
      <div class="form-group"><label>С</label>
        <input type="datetime-local" id="date_from" value="{{ default_from }}">
      </div>
      <div class="form-group"><label>По</label>
        <input type="datetime-local" id="date_to" value="{{ default_to }}">
      </div>
      <div class="form-group"><label>Записей (макс)</label>
        <input type="number" id="limit" value="1000" style="width:100px;">
      </div>
    </div>
    <div class="export-buttons">
      <a class="export-btn csv" id="btn-csv" href="#">Скачать CSV</a>
      <a class="export-btn excel" id="btn-excel" href="#">Скачать Excel</a>
      <button type="button" class="btn primary" id="btn-preview" style="padding:8px 20px;">Предпросмотр</button>
    </div>
  </form>
</div>

<div class="form-card" id="preview-area" style="overflow-x:auto;">
  <h2>Последние записи</h2>
  <div id="preview-content">
    {{ preview_html | safe }}
  </div>
</div>

<script>
// Export
function buildUrl(format) {
  var type = document.getElementById('data_type').value;
  var device = document.getElementById('device_filter').value;
  var from = document.getElementById('date_from').value;
  var to = document.getElementById('date_to').value;
  var limit = document.getElementById('limit').value;
  var url = '/api/export/' + type + '?format=' + format;
  if (device) url += '&device=' + encodeURIComponent(device);
  if (from) url += '&from=' + encodeURIComponent(from);
  if (to) url += '&to=' + encodeURIComponent(to);
  if (limit) url += '&limit=' + limit;
  return url;
}
document.getElementById('btn-csv').addEventListener('click', function(e) {
  e.preventDefault(); this.href = buildUrl('csv'); window.location = this.href;
});
document.getElementById('btn-excel').addEventListener('click', function(e) {
  e.preventDefault(); this.href = buildUrl('excel'); window.location = this.href;
});
document.getElementById('btn-preview').addEventListener('click', function() {
  fetch(buildUrl('html')).then(function(r){return r.text()}).then(function(h) {
    document.getElementById('preview-content').innerHTML = h;
  });
});

// Charts
var COLORS = ['#ff9800','#29b6f6','#66bb6a','#ab47bc','#ef5350','#26c6da','#ffca28','#8d6e63'];
var chartOpts = {
  responsive: true, maintainAspectRatio: false, animation: false,
  interaction: { mode: 'index', intersect: false },
  plugins: {
    legend: { labels: { color: '#90a4ae', boxWidth: 12, padding: 10 } },
    tooltip: { backgroundColor: '#1a2733', titleColor: '#4fc3f7', bodyColor: '#e0e0e0' }
  },
  scales: {
    x: {
      type: 'time',
      time: { tooltipFormat: 'dd.MM HH:mm', displayFormats: { minute: 'HH:mm', hour: 'HH:mm', day: 'dd.MM' } },
      ticks: { color: '#607d8b', maxTicksLimit: 10 },
      grid: { color: '#1a2733' }
    },
    y: { ticks: { color: '#607d8b' }, grid: { color: '#1a2733' } }
  }
};
var ctxT = document.getElementById('chartTemp').getContext('2d');
var ctxH = document.getElementById('chartHum').getContext('2d');
var ctxC = document.getElementById('chartCO2').getContext('2d');
var chartTemp = new Chart(ctxT, { type: 'line', data: { datasets: [] }, options: structuredClone(chartOpts) });
var chartHum = new Chart(ctxH, { type: 'line', data: { datasets: [] }, options: structuredClone(chartOpts) });
var chartCO2 = new Chart(ctxC, { type: 'line', data: { datasets: [] }, options: structuredClone(chartOpts) });
var refreshTimer = null;

function getSelectedDevices() {
  var devs = [];
  document.querySelectorAll('.chart-device:checked').forEach(function(cb){ devs.push(cb.value); });
  return devs;
}

function updateCharts() {
  var range = document.getElementById('chart_range').value;
  var devices = getSelectedDevices();
  var showT = document.getElementById('show_temp').checked;
  var showH = document.getElementById('show_hum').checked;
  var showC = document.getElementById('show_co2').checked;
  var url = '/api/chart_data?minutes=' + range;
  if (devices.length) url += '&devices=' + devices.join(',');

  fetch(url).then(function(r){return r.json()}).then(function(data) {
    var dsT = [], dsH = [], dsC = [];
    var i = 0;
    for (var dev in data) {
      var color = COLORS[i % COLORS.length];
      var points = data[dev];
      if (showT) {
        dsT.push({
          label: dev, borderColor: color, backgroundColor: color + '22',
          data: points.filter(function(p){return p.t!==null}).map(function(p){return {x:p.ts, y:p.t}}),
          borderWidth: 2, pointRadius: 0, tension: 0.3, fill: false
        });
      }
      if (showH) {
        dsH.push({
          label: dev, borderColor: color, backgroundColor: color + '22',
          data: points.filter(function(p){return p.h!==null}).map(function(p){return {x:p.ts, y:p.h}}),
          borderWidth: 2, pointRadius: 0, tension: 0.3, fill: false
        });
      }
      if (showC) {
        dsC.push({
          label: dev, borderColor: color, backgroundColor: color + '22',
          data: points.filter(function(p){return p.co2!==null}).map(function(p){return {x:p.ts, y:p.co2}}),
          borderWidth: 2, pointRadius: 0, tension: 0.3, fill: false
        });
      }
      i++;
    }
    chartTemp.data.datasets = dsT; chartTemp.update('none');
    chartHum.data.datasets = dsH; chartHum.update('none');
    chartCO2.data.datasets = dsC; chartCO2.update('none');
    // Show/hide chart containers
    chartTemp.canvas.parentElement.parentElement.style.display = showT && dsT.length ? '' : 'none';
    chartHum.canvas.parentElement.parentElement.style.display = showH && dsH.length ? '' : 'none';
    chartCO2.canvas.parentElement.parentElement.style.display = showC && dsC.length ? '' : 'none';
  });
}

function setupRefresh() {
  if (refreshTimer) clearInterval(refreshTimer);
  var sec = parseInt(document.getElementById('chart_refresh').value);
  if (sec > 0) refreshTimer = setInterval(updateCharts, sec * 1000);
}

document.getElementById('chart_range').addEventListener('change', updateCharts);
document.getElementById('chart_refresh').addEventListener('change', setupRefresh);
document.getElementById('show_temp').addEventListener('change', updateCharts);
document.getElementById('show_hum').addEventListener('change', updateCharts);
document.getElementById('show_co2').addEventListener('change', updateCharts);
document.querySelectorAll('.chart-device').forEach(function(cb){ cb.addEventListener('change', updateCharts); });

updateCharts();
setupRefresh();
</script>
"""


# ===================== HELPERS =====================

def format_uptime(s):
    if s < 60: return f"{s}с"
    if s < 3600: return f"{s // 60}м {s % 60}с"
    return f"{s // 3600}ч {(s % 3600) // 60}м"

def get_snapshot():
    with lock:
        s = {k: v for k, v in state.items()}
        s["relays"] = dict(state["relays"])
        s["sensors"] = {k: dict(v) for k, v in state["sensors"].items()}
    return s

def render_state():
    s = get_snapshot()
    return render_template_string(STATE_FRAGMENT, relays=s["relays"],
        arduino_online=s["arduino_online"], uptime_str=format_uptime(s["uptime"]),
        rssi=s["rssi"], names=get_relay_names(), groups=get_relay_groups(), sensors=s["sensors"],
        profiles=profiles, active_profiles=active_profiles)

def render_profiles_list():
    with profiles_lock:
        p_copy = [dict(p) for p in profiles]
    return render_template_string(PROFILES_LIST_FRAGMENT, profiles=p_copy,
        active_profiles=active_profiles, relay_names=get_relay_names())

def render_profiles_page():
    s = get_snapshot()
    with profiles_lock:
        p_copy = [dict(p) for p in profiles]
    list_html = render_template_string(PROFILES_LIST_FRAGMENT, profiles=p_copy,
        active_profiles=active_profiles, relay_names=get_relay_names())
    page = PROFILES_PAGE.replace('PROFILES_LIST_PLACEHOLDER', list_html)
    return render_template_string(page, arduino_online=s["arduino_online"],
        sensors=s["sensors"], relay_names=get_relay_names(),
        profiles=p_copy, active_profiles=active_profiles)

def render_rules_list():
    with rules_lock:
        rules_copy = [dict(r) for r in rules]
        active_copy = dict(rule_active)
    return render_template_string(RULES_LIST_FRAGMENT, rules=rules_copy,
        active_map=active_copy, relay_names=get_relay_names())

def render_rules_page():
    s = get_snapshot()
    sensor_list = sorted(s["sensors"].keys())
    with rules_lock:
        rules_copy = [dict(r) for r in rules]
        active_copy = dict(rule_active)
    list_html = render_template_string(RULES_LIST_FRAGMENT, rules=rules_copy,
        active_map=active_copy, relay_names=get_relay_names())
    page = RULES_PAGE.replace('RULES_LIST_PLACEHOLDER', list_html)
    with profiles_lock:
        p_list = [dict(p) for p in profiles]
    return render_template_string(page, arduino_online=s["arduino_online"],
        sensors=s["sensors"], sensor_list=sensor_list, relay_names=get_relay_names(),
        rules=rules_copy, active_map=active_copy, profiles_list=p_list)

def render_stats_page():
    s = get_snapshot()
    sensor_list = sorted(s["sensors"].keys())
    now = now_local()
    default_from = (now - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")
    default_to = now.strftime("%Y-%m-%dT%H:%M")
    # Preview last 20 sensor records
    preview_html = get_export_html("sensors", "", "", "", 20)
    return render_template_string(STATS_PAGE, arduino_online=s["arduino_online"],
        sensors=s["sensors"], sensor_list=sensor_list, relay_names=get_relay_names(),
        default_from=default_from, default_to=default_to, preview_html=preview_html)


def get_export_data(data_type, device, date_from, date_to, limit):
    """Query DB for export data."""
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    limit = min(int(limit or 10000), 50000)

    if data_type == "relays":
        query = "SELECT ts, relay_id, state FROM relay_log WHERE 1=1"
        params = []
        if device and device.startswith("relay_"):
            query += " AND relay_id = ?"
            params.append(int(device.replace("relay_", "")))
        if date_from:
            query += " AND ts >= ?"
            params.append(date_from.replace("T", " "))
        if date_to:
            query += " AND ts <= ?"
            params.append(date_to.replace("T", " "))
        query += " ORDER BY ts DESC LIMIT ?"
        params.append(limit)
        rows = conn.execute(query, params).fetchall()
        conn.close()
        headers = ["Время", "Реле", "Состояние"]
        data = []
        for r in rows:
            relay_name = get_relay_names().get(r["relay_id"], f"Реле {r['relay_id']}")
            data.append([r["ts"], relay_name, r["state"]])
        return headers, data
    else:
        query = "SELECT ts, device, temperature, humidity, co2 FROM sensor_log WHERE 1=1"
        params = []
        if device and not device.startswith("relay_"):
            query += " AND device = ?"
            params.append(device)
        if date_from:
            query += " AND ts >= ?"
            params.append(date_from.replace("T", " "))
        if date_to:
            query += " AND ts <= ?"
            params.append(date_to.replace("T", " "))
        query += " ORDER BY ts DESC LIMIT ?"
        params.append(limit)
        rows = conn.execute(query, params).fetchall()
        conn.close()
        headers = ["Время", "Датчик", "Температура", "Влажность", "CO2"]
        data = [[r["ts"], r["device"], r["temperature"], r["humidity"], r["co2"]] for r in rows]
        return headers, data

def get_export_html(data_type, device, date_from, date_to, limit):
    headers, data = get_export_data(data_type, device, date_from, date_to, limit)
    if not data:
        return "<p style='color:#90a4ae;'>Нет данных за выбранный период.</p>"
    html = '<table class="stats-table"><thead><tr>'
    for h in headers:
        html += f"<th>{h}</th>"
    html += "</tr></thead><tbody>"
    for row in data:
        html += "<tr>"
        for v in row:
            html += f"<td>{v if v is not None else '—'}</td>"
        html += "</tr>"
    html += "</tbody></table>"
    return html


# ===================== ROUTES =====================

@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        if request.form.get("password") == PASSWORD:
            session["logged_in"] = True
            return redirect("/")
        error = "Неверный пароль"
    return render_template_string(TEMPLATE, content=render_template_string(LOGIN_PAGE, error=error))

@app.route("/logout")
def logout():
    session.pop("logged_in", None)
    return redirect("/login")

@app.route("/")
@login_required
def index():
    return render_template_string(TEMPLATE, content=render_state())

@app.route("/state")
@login_required
def get_state_route():
    return render_state()

@app.route("/rules")
@login_required
def rules_page_route():
    return render_template_string(TEMPLATE, content=render_rules_page())

@app.route("/profiles")
@login_required
def profiles_page_route():
    return render_template_string(TEMPLATE, content=render_profiles_page())

@app.route("/stats")
@login_required
def stats_page_route():
    return render_template_string(TEMPLATE, content=render_stats_page())

@app.route("/api/profiles", methods=["POST"])
@login_required
def create_profile():
    f = request.form
    pre = [{"relay": i, "state": "ON"} for i in [14,12,11,9] if f.get(f"pre_{i}")]
    main = [{"relay": i, "state": "ON"} for i in [10,8,15,13] if f.get(f"main_{i}")]
    if not pre and not main:
        return render_profiles_list(), 400
    profile = {
        "id": str(uuid.uuid4())[:8],
        "name": f.get("name", "Без названия").strip(),
        "description": f.get("description", "").strip(),
        "pre": pre,
        "delay": int(f.get("delay", 10)),
        "delay_off": int(f.get("delay_off", 10)),
        "main": main,
    }
    with profiles_lock:
        profiles.append(profile)
        save_profiles()
    return render_profiles_list()

@app.route("/api/profile/<pid>/toggle", methods=["POST"])
@login_required
def toggle_profile(pid):
    toggle_profile_async(pid)
    time.sleep(0.3)
    return render_profiles_list()

@app.route("/api/profile/<pid>", methods=["DELETE"])
@login_required
def delete_profile(pid):
    with profiles_lock:
        active_profiles.pop(pid, None)
        profiles[:] = [p for p in profiles if p.get("id") != pid]
        save_profiles()
    return render_profiles_list()

@app.route("/relay/<int:num>/toggle", methods=["POST"])
@login_required
def toggle_relay(num):
    if 1 <= num <= 15:
        with lock:
            new_state = not state["relays"].get(num, False)
        mqtt_client.publish(f"greenhouse/relay/{num}/set", "ON" if new_state else "OFF")
    return "", 204

@app.route("/relay/all/on", methods=["POST"])
@login_required
def all_on():
    for i in range(1, 16):
        mqtt_client.publish(f"greenhouse/relay/{i}/set", "ON")
    return "", 204

@app.route("/relay/all/off", methods=["POST"])
@login_required
def all_off():
    for i in range(1, 16):
        mqtt_client.publish(f"greenhouse/relay/{i}/set", "OFF")
    return "", 204

@app.route("/group/<name>/on", methods=["POST"])
@login_required
def group_on(name):
    for rid in get_relay_groups().get(name, []):
        mqtt_client.publish(f"greenhouse/relay/{rid}/set", "ON")
    return "", 204

@app.route("/group/<name>/off", methods=["POST"])
@login_required
def group_off(name):
    for rid in get_relay_groups().get(name, []):
        mqtt_client.publish(f"greenhouse/relay/{rid}/set", "OFF")
    return "", 204

@app.route("/api/rules", methods=["POST"])
@login_required
def create_rule():
    f = request.form
    edit_id = f.get("edit_id", "").strip()
    conditions = []
    cond_device = f.get("cond_device", "").strip()
    if cond_device:
        conditions.append({
            "device": cond_device, "field": f.get("cond_field", "t"),
            "op": f.get("cond_op", ">"),
            "value": float(f.get("cond_value", 0)),
            "hysteresis": float(f.get("cond_hysteresis", 0)),
        })
    sched_start = f.get("sched_start", "").strip()
    sched_end = f.get("sched_end", "").strip()
    schedule = {"enabled": False}
    if sched_start and sched_end:
        sh, sm = map(int, sched_start.split(":"))
        eh, em = map(int, sched_end.split(":"))
        days = [i for i in range(7) if f.get(f"day_{i}")]
        schedule = {"enabled": True, "start_minutes": sh*60+sm, "end_minutes": eh*60+em,
                     "days": days if days else [0,1,2,3,4,5,6]}
    profile_id = f.get("profile_id", "").strip()
    actions = [{"relay": i, "command": "ON"} for i in range(1, 16) if f.get(f"relay_{i}")]
    if not actions and not profile_id:
        return render_rules_list(), 400

    rule = {
        "id": edit_id if edit_id else str(uuid.uuid4())[:8],
        "name": f.get("name", "Без названия").strip(), "enabled": True,
        "conditions": conditions, "schedule": schedule,
        "actions": actions, "reverse": "reverse" in f,
    }
    if profile_id:
        profile = find_profile(profile_id)
        rule["profile_id"] = profile_id
        rule["profile_name"] = profile["name"] if profile else profile_id
    with rules_lock:
        if edit_id:
            rules[:] = [r for r in rules if r.get("id") != edit_id]
        rules.append(rule)
        save_rules()
    return render_rules_list()

@app.route("/api/rules/<rule_id>/toggle", methods=["POST"])
@login_required
def toggle_rule(rule_id):
    with rules_lock:
        for r in rules:
            if r.get("id") == rule_id:
                r["enabled"] = not r.get("enabled", True)
                if not r["enabled"]:
                    rule_active.pop(rule_id, None)
                break
        save_rules()
    return render_rules_list()

@app.route("/api/rules/<rule_id>", methods=["DELETE"])
@login_required
def delete_rule(rule_id):
    with rules_lock:
        rules[:] = [r for r in rules if r.get("id") != rule_id]
        rule_active.pop(rule_id, None)
        save_rules()
    return render_rules_list()

@app.route("/api/chart_data")
@login_required
def chart_data():
    minutes = int(request.args.get("minutes", 360))
    devices = request.args.get("devices", "")
    device_list = [d.strip() for d in devices.split(",") if d.strip()] if devices else []

    now = now_local()
    since = (now - timedelta(minutes=minutes)).strftime("%Y-%m-%d %H:%M:%S")

    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    query = "SELECT ts, device, temperature, humidity, co2 FROM sensor_log WHERE ts >= ?"
    params = [since]
    if device_list:
        placeholders = ",".join("?" * len(device_list))
        query += f" AND device IN ({placeholders})"
        params.extend(device_list)
    query += " ORDER BY ts ASC LIMIT 10000"
    rows = conn.execute(query, params).fetchall()
    conn.close()

    result = {}
    for r in rows:
        dev = r["device"]
        if dev not in result:
            result[dev] = []
        result[dev].append({
            "ts": r["ts"],
            "t": r["temperature"],
            "h": r["humidity"],
            "co2": r["co2"],
        })
    return jsonify(result)

@app.route("/api/rename", methods=["POST"])
@login_required
def rename_item():
    data = request.get_json()
    item_type = data.get("type", "")
    item_id = data.get("id", "")
    new_name = data.get("name", "").strip()
    if not new_name or not item_id:
        return "", 400
    if item_type == "relay":
        custom_names.setdefault("relays", {})[str(item_id)] = new_name
    elif item_type == "group":
        custom_names.setdefault("groups", {})[str(item_id)] = new_name
    save_names()
    return "", 204

@app.route("/api/export/<data_type>")
@login_required
def export_data(data_type):
    device = request.args.get("device", "")
    date_from = request.args.get("from", "")
    date_to = request.args.get("to", "")
    limit = request.args.get("limit", "1000")
    fmt = request.args.get("format", "csv")

    if fmt == "html":
        return get_export_html(data_type, device, date_from, date_to, int(limit))

    headers, data = get_export_data(data_type, device, date_from, date_to, int(limit))

    if fmt == "excel":
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Данные"
            ws.append(headers)
            for row in data:
                ws.append(row)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return Response(buf.getvalue(),
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename=greenhouse_{data_type}.xlsx"})
        except ImportError:
            fmt = "csv"  # fallback

    # CSV
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(data)
    return Response(buf.getvalue(), mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=greenhouse_{data_type}.csv"})


if __name__ == "__main__":
    init_db()
    load_rules()
    load_names()
    load_profiles()
    start_mqtt()
    threading.Thread(target=rules_evaluation_loop, daemon=True).start()
    threading.Thread(target=history_logging_loop, daemon=True).start()
    app.run(host="0.0.0.0", port=8081)
